"""Documents without Office: DOCX, Markdown, HTML, plain text and PDF output."""

from __future__ import annotations

import html
import re
from html.parser import HTMLParser
from pathlib import Path

from ..registry import Options, Registry
from .data import _read_text

# ---------- DOCX reading ----------


def _iter_blocks(doc):
    """Paragraphs and tables in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def _heading_level(par) -> int:
    name = (par.style.name if par.style is not None else "") or ""
    m = re.match(r"Heading (\d)", name)
    if m:
        return int(m.group(1))
    if name == "Title":
        return 1
    return 0


def _is_list(par) -> bool:
    name = (par.style.name if par.style is not None else "") or ""
    if "List" in name:
        return True
    ppr = par._p.pPr
    return ppr is not None and ppr.numPr is not None


def _md_inline(par) -> str:
    out = []
    for run in par.runs:
        t = run.text
        if not t:
            continue
        if run.bold and t.strip():
            t = f"**{t}**"
        if run.italic and t.strip():
            t = f"*{t}*"
        out.append(t)
    return "".join(out) or par.text


def docx_to_markdown(src: Path) -> str:
    import docx

    doc = docx.Document(src)
    lines: list[str] = []
    for block in _iter_blocks(doc):
        if block.__class__.__name__ == "Table":
            rows = [[c.text.replace("\n", " ").replace("|", "\\|").strip() for c in r.cells] for r in block.rows]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("|" + " --- |" * width)
            lines.extend("| " + " | ".join(r) + " |" for r in rows[1:])
            lines.append("")
            continue
        text = _md_inline(block).strip()
        if not text:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        level = _heading_level(block)
        if level:
            lines.extend([f"{'#' * min(level, 6)} {block.text.strip()}", ""])
        elif _is_list(block):
            lines.append(f"- {text}")
        else:
            lines.extend([text, ""])
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


def docx_to_txt(src: Path, dst: Path, opts: Options) -> None:
    import docx

    doc = docx.Document(src)
    parts = []
    for block in _iter_blocks(doc):
        if block.__class__.__name__ == "Table":
            for r in block.rows:
                parts.append("\t".join(c.text for c in r.cells))
        else:
            parts.append(block.text)
    dst.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")


def docx_to_md(src: Path, dst: Path, opts: Options) -> None:
    dst.write_text(docx_to_markdown(src), encoding="utf-8")


# ---------- Markdown / HTML / text ----------

HTML_SHELL = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>body{{max-width:46rem;margin:2rem auto;padding:0 1rem;font:16px/1.6 system-ui,sans-serif}}
pre,code{{background:#f4f4f5;border-radius:4px}}pre{{padding:1rem;overflow:auto}}
table{{border-collapse:collapse}}td,th{{border:1px solid #ccc;padding:.3rem .6rem}}</style>
</head><body>
{body}
</body></html>
"""


def markdown_to_html_fragment(text: str) -> str:
    import markdown

    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"], output_format="html")


def md_to_html(src: Path, dst: Path, opts: Options) -> None:
    body = markdown_to_html_fragment(_read_text(src))
    dst.write_text(HTML_SHELL.format(title=html.escape(src.stem), body=body), encoding="utf-8")


def txt_to_html(src: Path, dst: Path, opts: Options) -> None:
    body = "<pre>" + html.escape(_read_text(src)) + "</pre>"
    dst.write_text(HTML_SHELL.format(title=html.escape(src.stem), body=body), encoding="utf-8")


class _TextExtractor(HTMLParser):
    BLOCK = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "pre",
        "blockquote",
        "table",
        "ul",
        "ol",
        "header",
        "footer",
    }
    SKIP = {"script", "style", "noscript", "template", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")
        if tag == "li":
            self.parts.append("- ")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(re.sub(r"[ \t\r\n]+", " ", data))

    def text(self) -> str:
        t = "".join(self.parts)
        t = "\n".join(line.strip() for line in t.splitlines())
        return re.sub(r"\n{3,}", "\n\n", t).strip() + "\n"


def html_to_text(markup: str) -> str:
    p = _TextExtractor()
    p.feed(markup)
    p.close()
    return p.text()


def html_to_txt(src: Path, dst: Path, opts: Options) -> None:
    dst.write_text(html_to_text(_read_text(src)), encoding="utf-8")


def md_to_txt(src: Path, dst: Path, opts: Options) -> None:
    dst.write_text(html_to_text(markdown_to_html_fragment(_read_text(src))), encoding="utf-8")


def txt_to_md(src: Path, dst: Path, opts: Options) -> None:
    text = _read_text(src)
    # escape characters that would otherwise turn plain text into markup
    escaped = re.sub(r"(?m)^(\s*)([#>*+-]|\d+\.)(\s)", r"\1\\\2\3", text)
    dst.write_text(escaped, encoding="utf-8")


# ---------- DOCX writing ----------


def _add_inline(par, text: str) -> None:
    for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            par.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = par.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            par.add_run(token[1:-1]).italic = True
        else:
            par.add_run(token)


def markdown_to_docx(text: str, dst: Path) -> None:
    import docx

    doc = docx.Document()
    in_code = False
    code: list[str] = []
    table: list[list[str]] = []

    def flush_table():
        if not table:
            return
        rows = [r for r in table if not all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in r)]
        width = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=width)
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                t.cell(i, j).text = c.strip()
        table.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph("\n".join(code))
                for run in p.runs:
                    run.font.name = "Consolas"
                code.clear()
            in_code = not in_code
            continue
        if in_code:
            code.append(line)
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table.append(line.strip().strip("|").split("|"))
            continue
        flush_table()
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif re.match(r"^\s*[-*+]\s+", line):
            _add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*+]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif line.strip():
            _add_inline(doc.add_paragraph(), line.strip())
    flush_table()
    if code:
        doc.add_paragraph("\n".join(code))
    doc.save(dst)


def md_to_docx(src: Path, dst: Path, opts: Options) -> None:
    markdown_to_docx(_read_text(src), dst)


def txt_to_docx(src: Path, dst: Path, opts: Options) -> None:
    import docx

    doc = docx.Document()
    for para in re.split(r"\n\s*\n", _read_text(src)):
        if para.strip():
            doc.add_paragraph(para.strip("\n"))
    doc.save(dst)


# ---------- PDF writing (reportlab) ----------

_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "C:/Windows/Fonts/arial.ttf",
)


def _font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path in _FONT_CANDIDATES:
        if Path(path).exists():
            try:
                if "Body" not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont("Body", path))
                return "Body"
            except Exception:  # noqa: BLE001, S112 - try the next font  # pragma: no cover
                continue
    return "Helvetica"


def text_to_pdf(text: str, dst: Path, *, title: str = "", markdown_mode: bool = False) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer

    font = _font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=font, fontSize=10.5, leading=14.5)
    mono = ParagraphStyle("mono", parent=body, fontName="Courier", fontSize=9, leading=11.5)
    heads = {
        i: ParagraphStyle(f"h{i}", parent=body, fontSize=20 - i * 2, leading=26 - i * 2, spaceBefore=8, spaceAfter=4)
        for i in range(1, 7)
    }
    story = []
    if markdown_mode:
        in_code, code = False, []
        for line in text.splitlines():
            if line.strip().startswith("```"):
                if in_code:
                    story.append(Preformatted("\n".join(code), mono))
                    code = []
                in_code = not in_code
                continue
            if in_code:
                code.append(line)
                continue
            m = re.match(r"^(#{1,6})\s+(.*)", line)
            esc = html.escape(line.strip())
            esc = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", esc)
            esc = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", esc)
            if m:
                story.append(Paragraph(html.escape(m.group(2)), heads[len(m.group(1))]))
            elif re.match(r"^\s*[-*+]\s+", line):
                story.append(Paragraph("• " + esc[2:].strip(), body))
            elif line.strip():
                story.append(Paragraph(esc, body))
            else:
                story.append(Spacer(1, 4))
        if code:
            story.append(Preformatted("\n".join(code), mono))
    else:
        for para in re.split(r"\n\s*\n", text):
            if para.strip():
                story.append(Paragraph(html.escape(para).replace("\n", "<br/>"), body))
                story.append(Spacer(1, 4))
    if not story:
        story.append(Spacer(1, 1))
    doc = SimpleDocTemplate(
        str(dst),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
    )
    doc.build(story)


def txt_to_pdf(src: Path, dst: Path, opts: Options) -> None:
    text_to_pdf(_read_text(src), dst, title=src.stem)


def md_to_pdf(src: Path, dst: Path, opts: Options) -> None:
    text_to_pdf(_read_text(src), dst, title=src.stem, markdown_mode=True)


def register(reg: Registry) -> None:
    reg.add("docx", "txt", docx_to_txt, engine="python-docx")
    reg.add("docx", "md", docx_to_md, engine="python-docx")
    reg.add("md", "html", md_to_html, engine="markdown")
    reg.add("md", "txt", md_to_txt, engine="markdown")
    reg.add("md", "docx", md_to_docx, engine="python-docx", cost=1.5)
    reg.add("md", "pdf", md_to_pdf, engine="reportlab", cost=1.4)
    reg.add("txt", "html", txt_to_html, engine="builtin")
    reg.add("txt", "md", txt_to_md, engine="builtin")
    reg.add("txt", "docx", txt_to_docx, engine="python-docx")
    reg.add("txt", "pdf", txt_to_pdf, engine="reportlab", cost=1.5)
    reg.add("html", "txt", html_to_txt, engine="builtin")
