"""Every built-in converter family, end to end through the pipeline."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from PIL import Image

from convertly import Options, convert_file
from convertly.converters.data import FORMATS as DATA_FORMATS
from convertly.converters.pdf import parse_pages

# ---------- images ----------


@pytest.mark.parametrize("fmt", ["jpg", "webp", "bmp", "gif", "tiff", "ico", "tga", "ppm", "pcx", "pdf"])
def test_png_to_everything(png, fmt):
    out = convert_file(png, fmt).outputs[0]
    assert out.suffix == f".{fmt}" and out.stat().st_size > 0
    if fmt != "pdf":
        with Image.open(out) as im:
            im.load()


def test_alpha_flattened_on_white_for_jpg(tmp_path):
    p = tmp_path / "t.png"
    Image.new("RGBA", (4, 4), (0, 0, 0, 0)).save(p)
    out = convert_file(p, "jpg").outputs[0]
    with Image.open(out) as im:
        assert im.getpixel((1, 1))[0] > 240


def test_resize_keeps_aspect(png):
    out = convert_file(png, "webp", options=Options(width=32)).outputs[0]
    with Image.open(out) as im:
        assert im.size == (32, 20)


def test_resize_box(png):
    out = convert_file(png, "jpg", options=Options(width=100, height=10)).outputs[0]
    with Image.open(out) as im:
        assert im.size == (16, 10)


def test_grayscale(png):
    out = convert_file(png, "jpg", options=Options(grayscale=True)).outputs[0]
    with Image.open(out) as im:
        assert im.mode == "L"


def test_quality_changes_size(tmp_path):
    import random

    random.seed(1)
    img = Image.new("RGB", (200, 200))
    img.putdata([(random.randrange(256), random.randrange(256), random.randrange(256)) for _ in range(40000)])
    p = tmp_path / "noise.png"
    img.save(p)
    lo = convert_file(p, "jpg", options=Options(quality=10)).outputs[0].stat().st_size
    hi = convert_file(p, "jpg", options=Options(quality=95)).outputs[0].stat().st_size
    assert lo < hi


def test_gps_exif_is_stripped(jpg_with_gps):
    out = convert_file(jpg_with_gps, "webp").outputs[0]
    with Image.open(out) as im:
        exif = im.getexif()
        assert 0x8825 not in exif and 0x010F not in exif


def test_animated_gif_to_webp_keeps_frames(animated_gif):
    out = convert_file(animated_gif, "webp").outputs[0]
    with Image.open(out) as im:
        assert getattr(im, "n_frames", 1) == 3


def test_animated_gif_to_pdf_multipage(animated_gif):
    out = convert_file(animated_gif, "pdf").outputs[0]
    from pypdf import PdfReader

    assert len(PdfReader(out).pages) == 3


def test_decompression_bomb_rejected(tmp_path, monkeypatch):
    import convertly.converters.images as images

    p = tmp_path / "big.png"
    Image.new("L", (2000, 2000)).save(p)
    monkeypatch.setattr(images, "MAX_PIXELS", 1_000_000)
    with pytest.raises(Exception, match="(?i)decompression|exceeds"):
        convert_file(p, "jpg")


# ---------- PDF ----------


@pytest.mark.parametrize(
    ("spec", "want"),
    [(None, [0, 1, 2]), ("2", [1]), ("1-2", [0, 1]), ("2-", [1, 2]), ("1,3", [0, 2]), ("1-99", [0, 1, 2])],
)
def test_parse_pages(spec, want):
    assert parse_pages(spec, 3) == want


@pytest.mark.parametrize("spec", ["0", "3-1", "9", "x"])
def test_parse_pages_invalid(spec):
    with pytest.raises(ValueError):
        parse_pages(spec, 3)


def test_pdf_to_txt(pdf3):
    out = convert_file(pdf3, "txt").outputs[0]
    text = out.read_text()
    assert "Hello page 1" in text and "Hello page 3" in text


def test_pdf_to_png_one_file_per_page(pdf3):
    r = convert_file(pdf3, "png", options=Options(dpi=72))
    assert [p.name for p in r.outputs] == ["doc-p1.png", "doc-p2.png", "doc-p3.png"]


def test_pdf_to_jpg_single_page(pdf3):
    r = convert_file(pdf3, "jpg", options=Options(pages="2", dpi=72))
    assert [p.name for p in r.outputs] == ["doc.jpg"]


def test_pdf_split(pdf3):
    r = convert_file(pdf3, "pages")
    assert len(r.outputs) == 3 and all(p.suffix == ".pdf" for p in r.outputs)


def test_pdf_merge_and_compress(pdf3, tmp_path):
    from pypdf import PdfReader

    from convertly.converters.pdf import compress_pdf, merge_pdfs

    merged = tmp_path / "m.pdf"
    merge_pdfs([pdf3, pdf3], merged)
    assert len(PdfReader(merged).pages) == 6
    small = tmp_path / "c.pdf"
    compress_pdf(merged, small, Options())
    assert len(PdfReader(small).pages) == 6
    assert small.stat().st_size <= merged.stat().st_size


def test_encrypted_pdf_reports_password(pdf3, tmp_path):
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter()
    for p in PdfReader(pdf3).pages:
        w.add_page(p)
    w.encrypt("secret")
    locked = tmp_path / "locked.pdf"
    with open(locked, "wb") as fh:
        w.write(fh)
    with pytest.raises(Exception, match="(?i)password"):
        convert_file(locked, "txt")


# ---------- data ----------

SAMPLE = {"people": [{"name": "Ana", "age": 31, "tags": ["a", "b"]}, {"name": "Bo", "age": 25, "tags": []}]}


def _write_sample(tmp_path: Path) -> Path:
    p = tmp_path / "sample.json"
    p.write_text(json.dumps(SAMPLE))
    return p


@pytest.mark.parametrize("fmt", [f for f in DATA_FORMATS if f != "json"])
def test_json_to_each_data_format_and_back(tmp_path, fmt):
    src = _write_sample(tmp_path)
    out = convert_file(src, fmt).outputs[0]
    back = convert_file(out, "json", out_dir=tmp_path / "back").outputs[0]
    data = json.loads(back.read_text())
    if fmt in ("json", "yaml", "toml", "jsonl"):
        flat = data if fmt != "jsonl" else data[0]
        assert "Ana" in json.dumps(flat)
    else:
        assert "Ana" in json.dumps(data)


def test_csv_rows(tmp_path):
    src = _write_sample(tmp_path)
    out = convert_file(src, "csv").outputs[0]
    rows = list(csv.DictReader(out.open(encoding="utf-8")))
    assert rows[0]["name"] == "Ana" and rows[0]["tags.0"] == "a" and rows[1]["age"] == "25"


def test_csv_formula_injection_neutralised(tmp_path):
    src = tmp_path / "evil.json"
    src.write_text(json.dumps([{"cmd": '=HYPERLINK("http://x")', "n": "-5", "at": "@SUM(A1)"}]))
    rows = list(csv.DictReader(convert_file(src, "csv").outputs[0].open(encoding="utf-8")))
    assert rows[0]["cmd"].startswith("'=") and rows[0]["at"].startswith("'@") and rows[0]["n"] == "-5"


def test_xml_billion_laughs_blocked(tmp_path):
    p = tmp_path / "bomb.xml"
    p.write_text('<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]><r>&b;</r>')
    with pytest.raises(Exception, match="(?i)entit"):
        convert_file(p, "json")


def test_xml_external_entity_blocked(tmp_path):
    p = tmp_path / "xxe.xml"
    p.write_text('<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]><r>&x;</r>')
    with pytest.raises(Exception):
        convert_file(p, "json")


def test_xml_roundtrip_attributes(tmp_path):
    p = tmp_path / "a.xml"
    p.write_text('<cat id="7"><item>one</item><item>two</item><name lang="en">Tom</name></cat>')
    data = json.loads(convert_file(p, "json").outputs[0].read_text())
    assert data == {"cat": {"@id": "7", "item": ["one", "two"], "name": {"@lang": "en", "#text": "Tom"}}}
    back = convert_file(
        convert_file(p, "json", out_dir=tmp_path / "j").outputs[0], "xml", out_dir=tmp_path / "x"
    ).outputs[0]
    assert 'id="7"' in back.read_text() and "<item>two</item>" in back.read_text()


def test_yaml_unsafe_tags_rejected(tmp_path):
    p = tmp_path / "evil.yaml"
    p.write_text("!!python/object/apply:os.system ['echo pwned']\n")
    with pytest.raises(Exception):
        convert_file(p, "json")


def test_toml_drops_nulls_and_wraps_lists(tmp_path):
    p = tmp_path / "l.json"
    p.write_text(json.dumps([{"a": 1, "b": None}]))
    text = convert_file(p, "toml").outputs[0].read_text()
    assert "items" in text and "a = 1" in text and "b =" not in text


def test_utf16_and_bom_input(tmp_path):
    p = tmp_path / "u.csv"
    p.write_bytes("name,city\nJosé,Zürich\n".encode("utf-16"))
    data = json.loads(convert_file(p, "json").outputs[0].read_text())
    assert data == [{"name": "José", "city": "Zürich"}]


# ---------- spreadsheets ----------


def test_csv_xlsx_roundtrip_types_and_safety(tmp_path):
    import openpyxl

    p = tmp_path / "t.csv"
    p.write_text("id,zip,price,formula\n1,02134,9.5,=1+1\n2,10001,3,=cmd|' /C calc'!A0\n")
    x = convert_file(p, "xlsx").outputs[0]
    ws = openpyxl.load_workbook(x).active
    assert ws["A2"].value == 1 and ws["B2"].value == "02134" and ws["C2"].value == 9.5
    assert ws["D2"].data_type == "s" and ws["D3"].data_type == "s"
    back = convert_file(x, "csv", out_dir=tmp_path / "b").outputs[0]
    rows = list(csv.DictReader(back.open(encoding="utf-8")))
    assert rows[0]["zip"] == "02134"


def test_xlsx_multi_sheet_to_json_and_sheet_option(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.active.append(["k", "v"])
    wb.active.append(["x", 1])
    ws = wb.create_sheet("B")
    ws.append(["k"])
    ws.append(["y"])
    p = tmp_path / "multi.xlsx"
    wb.save(p)
    data = json.loads(convert_file(p, "json").outputs[0].read_text())
    assert data == {"A": [{"k": "x", "v": 1}], "B": [{"k": "y"}]}
    one = convert_file(p, "csv", options=Options(sheet="B")).outputs[0].read_text()
    assert "y" in one and "x" not in one
    with pytest.raises(ValueError, match="not found"):
        convert_file(p, "csv", options=Options(sheet="nope"))


# ---------- documents ----------

MD = """# Title

Some **bold** and *italic* text.

## Section

- one
- two

| a | b |
|---|---|
| 1 | 2 |
"""


def test_md_docx_md_roundtrip(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text(MD)
    d = convert_file(p, "docx").outputs[0]
    back = convert_file(d, "md", out_dir=tmp_path / "b").outputs[0].read_text()
    assert "# Title" in back and "## Section" in back and "- one" in back and "**bold**" in back
    assert "| a | b |" in back


def test_md_to_html_and_txt(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text(MD)
    html = convert_file(p, "html").outputs[0].read_text()
    assert "<h1>Title</h1>" in html and "<table>" in html and "<strong>bold</strong>" in html
    txt = convert_file(p, "txt").outputs[0].read_text()
    assert "Title" in txt and "**" not in txt


def test_html_to_txt_skips_scripts(tmp_path):
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head><style>x{}</style></head><body><h1>Hi</h1><script>alert(1)</script>"
        "<p>A &amp; B</p><ul><li>x</li></ul></body></html>"
    )
    txt = convert_file(p, "txt").outputs[0].read_text()
    assert "Hi" in txt and "A & B" in txt and "alert" not in txt and "- x" in txt


def test_txt_to_html_escapes(tmp_path):
    p = tmp_path / "t.txt"
    p.write_text("<script>alert(1)</script>")
    html = convert_file(p, "html").outputs[0].read_text()
    assert "<script>alert" not in html and "&lt;script&gt;" in html


@pytest.mark.parametrize("src_fmt", ["txt", "md"])
def test_text_to_pdf_unicode(tmp_path, src_fmt):
    from pypdf import PdfReader

    p = tmp_path / f"u.{src_fmt}"
    p.write_text("Héllo wörld — ünïcode\n\nSecond <para> & more\n" + "line\n" * 120)
    out = convert_file(p, "pdf").outputs[0]
    r = PdfReader(out)
    assert len(r.pages) >= 2 and "Héllo" in r.pages[0].extract_text()


def test_txt_to_docx_and_md(tmp_path):
    import docx

    p = tmp_path / "t.txt"
    p.write_text("# not a heading\n\nPara two\n")
    d = docx.Document(convert_file(p, "docx").outputs[0])
    assert [x.text for x in d.paragraphs] == ["# not a heading", "Para two"]
    md = convert_file(p, "md").outputs[0].read_text()
    assert md.startswith("\\# not a heading")


# ---------- archives ----------


def _make_zip(tmp_path: Path) -> Path:
    import zipfile

    p = tmp_path / "bundle.zip"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("a.txt", "alpha")
        zf.writestr("sub/b.txt", "beta")
    return p


@pytest.mark.parametrize("fmt", ["tar", "tar.gz", "tar.bz2", "tar.xz"])
def test_zip_to_tar_and_back(tmp_path, fmt):
    import zipfile

    z = _make_zip(tmp_path)
    t = convert_file(z, fmt).outputs[0]
    assert t.name == f"bundle.{fmt}"
    back = convert_file(t, "zip", out_dir=tmp_path / "back").outputs[0]
    with zipfile.ZipFile(back) as zf:
        assert sorted(zf.namelist()) == ["a.txt", "sub/b.txt"] and zf.read("sub/b.txt") == b"beta"


def test_extract_to_folder(tmp_path):
    z = _make_zip(tmp_path)
    out = convert_file(z, "dir").outputs[0]
    assert out.is_dir() and (out / "sub" / "b.txt").read_text() == "beta"
    again = convert_file(z, "dir").outputs[0]
    assert again.name == "bundle (1)"


def test_pack_single_file_to_zip(png):
    import zipfile

    out = convert_file(png, "zip").outputs[0]
    with zipfile.ZipFile(out) as zf:
        assert zf.namelist() == ["photo.png"]


def test_zip_slip_rejected(tmp_path):
    import zipfile

    p = tmp_path / "evil.zip"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("../../escape.txt", "x")
    with pytest.raises(Exception, match="escapes"):
        convert_file(p, "dir", out_dir=tmp_path / "o")
    assert not (tmp_path / "escape.txt").exists()


def test_tar_absolute_and_symlink_members(tmp_path):
    import io
    import tarfile

    p = tmp_path / "evil.tar"
    with tarfile.open(p, "w") as tf:
        link = tarfile.TarInfo("link")
        link.type = tarfile.SYMTYPE
        link.linkname = "/etc/passwd"
        tf.addfile(link)
        data = b"ok"
        info = tarfile.TarInfo("fine.txt")
        info.size = len(data)
        tf.addfile(info, io.BytesIO(data))
    out = convert_file(p, "dir").outputs[0]
    assert (out / "fine.txt").read_text() == "ok" and not (out / "link").exists()

    p2 = tmp_path / "abs.tar"
    with tarfile.open(p2, "w") as tf:
        info = tarfile.TarInfo("/tmp/abs.txt")
        info.size = 1
        tf.addfile(info, io.BytesIO(b"x"))
    with pytest.raises(Exception, match="escapes"):
        convert_file(p2, "dir")


def test_zip_bomb_ratio_rejected(tmp_path, monkeypatch):
    import zipfile

    import convertly.safety as safety

    p = tmp_path / "bomb.zip"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("zeros.bin", b"\0" * (60 * 1024**2))
    with pytest.raises(Exception, match="(?i)unsafe size"):
        convert_file(p, "dir")
    monkeypatch.setattr(safety, "MAX_ARCHIVE_MEMBERS", 0)
    with pytest.raises(Exception, match="too many"):
        convert_file(_make_zip(tmp_path), "dir")


def test_heic_input_when_pillow_heif_installed(tmp_path):
    pillow_heif = pytest.importorskip("pillow_heif")
    pillow_heif.register_heif_opener()
    p = tmp_path / "phone.heic"
    Image.new("RGB", (40, 30), (200, 10, 10)).save(p)
    out = convert_file(p, "jpg").outputs[0]
    with Image.open(out) as im:
        assert im.size == (40, 30)


RICH_MD = """# Guide

1. first
2. second

* star bullet with `code` and **bold**

```
def hi():
    return "<tag> & stuff"
```

| x | y |
|---|---|
| 1 | 2 |

### Small
Text with *italic*.
"""


@pytest.mark.parametrize("fmt", ["docx", "pdf", "html", "txt"])
def test_rich_markdown_everywhere(tmp_path, fmt):
    p = tmp_path / "rich.md"
    p.write_text(RICH_MD)
    out = convert_file(p, fmt).outputs[0]
    assert out.stat().st_size > 0
    if fmt == "docx":
        import docx

        from convertly import Options
        from convertly.converters.documents import md_to_docx

        assert "stuff" in "\n".join(x.text for x in docx.Document(out).paragraphs)  # whichever engine ran
        builtin = tmp_path / "builtin.docx"
        md_to_docx(p, builtin, Options())  # the no-Pandoc path
        d = docx.Document(builtin)
        styles = [x.style.name for x in d.paragraphs]
        texts = "\n".join(x.text for x in d.paragraphs)
        assert "List Number" in styles and "List Bullet" in styles and "Heading 3" in styles
        assert 'return "<tag> & stuff"' in texts and len(d.tables) == 1
    if fmt == "pdf":
        from pypdf import PdfReader

        text = PdfReader(out).pages[0].extract_text()
        assert "Guide" in text and "stuff" in text


def test_docx_tables_and_text(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Intro")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "h1", "h|2", "a", "b"
    d.add_paragraph("item", style="List Bullet")
    p = tmp_path / "t.docx"
    d.save(p)
    txt = convert_file(p, "txt").outputs[0].read_text()
    assert "h1\th|2" in txt and "item" in txt
    md = convert_file(p, "md").outputs[0].read_text()
    assert "| h1 | h\\|2 |" in md and "- item" in md
